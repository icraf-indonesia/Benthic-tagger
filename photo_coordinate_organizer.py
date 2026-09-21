"""Photo and Coordinate Organizer for Marine Transect Surveys.

Extracts coordinates from CSV/KML files, organizes/renames survey photos based on
Day, Transect, and Point IDs, extracts image EXIF metadata, and generates an Excel
workbook matching the exact layout, column structure, styling, and formulas of Example.xlsx.

Supports any UTM Zone across Indonesia (e.g. Zone 48S/48M & 49S/49M in Bangka Belitung,
Zone 51S/51M in Sulawesi) with automatic zone detection or user-specified zone enforcement.
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    from PIL import Image, ExifTags
except ImportError as error:
    raise SystemExit("Pillow is required: pip install Pillow") from error

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as error:
    raise SystemExit("openpyxl is required: pip install openpyxl") from error


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}


def get_utm_band(lat: float) -> str:
    """Get standard MGRS/UTM latitude band letter.
    
    For latitudes between -8° and 0° (e.g. Bangka Belitung, Central Sulawesi),
    the band letter is 'M'.
    """
    bands = "CDEFGHJKLMNPQRSTUVWX"
    if -80 <= lat <= 84:
        return bands[int((lat + 80) / 8)]
    return "M" if lat < 0 else "N"


def get_utm_columns(zone: int = 51, band: str = "M", format_type: str = "MGRS") -> Tuple[str, str]:
    """Generate UTM Easting and Northing column names.
    
    - 'MGRS' (default, matches Example.xlsx): X_UTM51M, X_UTM48M, X_UTM49M
    - 'HEMISPHERE': X_UTM51S, X_UTM48S, X_UTM49S
    """
    if format_type.upper() == "HEMISPHERE":
        suffix = "S" if band in "CDEFGHJKLM" else "N"
    else:
        suffix = band
    return f"X_UTM{zone}{suffix}", f"Y_UTM{zone}{suffix}"


def build_column_list(zone: int = 51, band: str = "M", format_type: str = "MGRS") -> List[str]:
    """Generate the standard 27 survey columns matching Example.xlsx with dynamic UTM headers."""
    col_x, col_y = get_utm_columns(zone, band, format_type)
    return [
        "ID",                        # Col A (1): Row index (1, 2, ...)
        "Team",                      # Col B (2): Survey team identifier ('A')
        "ID Transek",                # Col C (3): Transect ID ('T1')
        "ID Titik",                  # Col D (4): Point ID ('D1T1-1')
        "File name",                 # Col E (5): Renamed file name ('D1T1 (1).JPG')
        "Tanggal",                   # Col F (6): Survey date from CSV
        "Jam",                       # Col G (7): Survey time from CSV
        "Variasi Habitat",           # Col H (8): Habitat variation ('?')
        "Kedalaman",                 # Col I (9): Depth/elevation ('?')
        "Range (m)",                 # Col J (10): Range distance (0)
        "Azimuth",                   # Col K (11): Direction/bearing (90)
        "Latitude",                  # Col L (12): Latitude in decimal degrees
        "Longitude",                 # Col M (13): Longitude in decimal degrees
        "Lokasi",                    # Col N (14): Location name ('Tandaigi')
        col_x,                       # Col O (15): Dynamic UTM Easting (e.g. X_UTM51M or X_UTM48S)
        col_y,                       # Col P (16): Dynamic UTM Northing (e.g. Y_UTM51M or Y_UTM48S)
        "Photo",                     # Col Q (17): In-cell image formula =IMAGE(R{row})
        "URL",                       # Col R (18): Image URL or file path
        "Annotation/Interpretation", # Col S (19): Left empty for human annotator
        "FileSize",                  # Col T (20): Photo file size ('5.9 MB')
        "Model",                     # Col U (21): Camera model ('V50 X Cube')
        "DateOriginal",              # Col V (22): Photo capture date from EXIF
        "TimeOriginal",              # Col W (23): Photo capture time from EXIF
        "ShutterSpeed",              # Col X (24): Exposure time formatted as fraction ('1/225')
        "Aperture",                  # Col Y (25): F-number formatted string ('2.8')
        "ISO",                       # Col Z (26): ISO speed rating (100)
        "WhiteBalance",              # Col AA (27): White balance mode ('Auto')
    ]


@dataclass
class CoordinateRecord:
    source_file: str
    location: str
    day: str
    title: str
    point_id: Optional[int]
    latitude: float
    longitude: float
    utm_x: float
    utm_y: float
    utm_zone: int
    utm_band: str
    date_created: Optional[str] = None
    survey_date: Optional[str] = None
    survey_time: Optional[str] = None
    elevation: Optional[str] = None


@dataclass
class PhotoMetadata:
    source_path: Path
    original_filename: str
    day: str
    transect: str
    point_id: Optional[int]
    sequence: Optional[int]
    renamed_filename: str
    file_size_str: str
    camera_model: Optional[str]
    date_original: Optional[str]
    time_original: Optional[str]
    shutter_speed: Optional[str]
    aperture: Optional[str]
    iso: Optional[int]
    white_balance: Optional[str]


def latlon_to_utm(
    lat: float,
    lon: float,
    zone: Optional[int] = None,
    south: Optional[bool] = None,
) -> Tuple[float, float, int, str]:
    """Convert WGS84 Latitude and Longitude to UTM projected coordinates.

    Automatically calculates UTM Zone from Longitude if not provided:
    Zone = int((lon + 180) / 6) + 1.
    - Bangka (~106°E): Zone 48
    - Belitung (~107.5° - 108.5°E): Zone 48 or 49
    - Central Sulawesi (~120°E): Zone 51

    Uses the standard Transverse Mercator projection on the WGS84 ellipsoid.
    Matches Example.xlsx coordinates down to the centimeter (e.g. 173339.73, 9931469.70).
    """
    if zone is None:
        zone = int((lon + 180) / 6) + 1
    if south is None:
        south = lat < 0

    band = get_utm_band(lat)

    a = 6378137.0  # WGS84 semi-major axis
    f = 1 / 298.257223563  # WGS84 flattening
    b = a * (1 - f)
    e = math.sqrt(1 - (b / a) ** 2)
    e_prime_sq = (e ** 2) / (1 - e ** 2)
    k0 = 0.9996

    lat_rad = math.radians(lat)
    lon_rad = math.radians(lon)
    lon0 = (zone - 1) * 6 - 180 + 3
    lon0_rad = math.radians(lon0)

    N = a / math.sqrt(1 - e ** 2 * math.sin(lat_rad) ** 2)
    T = math.tan(lat_rad) ** 2
    C = e_prime_sq * math.cos(lat_rad) ** 2
    A = math.cos(lat_rad) * (lon_rad - lon0_rad)

    M = a * (
        (1 - e ** 2 / 4 - 3 * e ** 4 / 64 - 5 * e ** 6 / 256) * lat_rad
        - (3 * e ** 2 / 8 + 3 * e ** 4 / 32 + 45 * e ** 6 / 1024) * math.sin(2 * lat_rad)
        + (15 * e ** 4 / 256 + 45 * e ** 6 / 1024) * math.sin(4 * lat_rad)
        - (35 * e ** 6 / 3072) * math.sin(6 * lat_rad)
    )

    x = (
        k0
        * N
        * (
            A
            + (1 - T + C) * A ** 3 / 6
            + (5 - 18 * T + T ** 2 + 72 * C - 58 * e_prime_sq) * A ** 5 / 120
        )
        + 500000.0
    )

    y = k0 * (
        M
        + N
        * math.tan(lat_rad)
        * (
            A ** 2 / 2
            + (5 - T + 9 * C + 4 * C ** 2) * A ** 4 / 24
            + (61 - 58 * T + T ** 2 + 600 * C - 330 * e_prime_sq) * A ** 6 / 720
        )
    )

    if south or lat < 0:
        y += 10000000.0

    return round(x, 2), round(y, 2), zone, band


def extract_number_from_text(text: str | None) -> Optional[int]:
    """Extract point number from text like 'Tanda tempat 1', 'Titik 2', 'Point 14'.
    Avoids matching transect markers such as 'Transek 2'.
    """
    if not text:
        return None
    match = re.search(r"(?:Tanda\s+tempat|Titik|Point|P)\s*(\d+)\b", text, re.IGNORECASE)
    if match:
        return int(match.group(1))
    if re.search(r"\bTransek\b", text, re.IGNORECASE):
        return None
    match = re.search(r"\b(\d+)\b", text)
    return int(match.group(1)) if match else None


def parse_coord_filename(filename: str) -> Tuple[str, str]:
    """Extract location name and day identifier from filename (e.g. 'Tandaigi_D1.csv')."""
    stem = Path(filename).stem
    match = re.search(r"^(.*?)[_-]?(D\d+)", stem, re.IGNORECASE)
    if match:
        location = match.group(1).rstrip("_- ") or "Tandaigi"
        day = match.group(2).upper()
        return location, day
    return stem, "D1"


def load_coordinates(
    coords_dir: str | Path,
    target_zone: Optional[int] = None,
) -> Dict[Tuple[str, int], CoordinateRecord]:
    """Load coordinate CSV files from coords_dir and return a dictionary keyed by (day, point_id).

    If target_zone is provided (e.g. 48 or 49 for Bangka Belitung), all coordinates
    are projected into that uniform UTM zone. Otherwise, each coordinate auto-detects its zone.
    """
    coords_path = Path(coords_dir)
    records: Dict[Tuple[str, int], CoordinateRecord] = {}

    if not coords_path.exists():
        return records

    csv_files = sorted([
        f for f in coords_path.iterdir()
        if f.is_file() and f.suffix.lower() == ".csv"
    ])

    for csv_file in csv_files:
        location, day = parse_coord_filename(csv_file.name)
        with open(csv_file, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                try:
                    title = row.get("Title") or row.get("Name") or ""
                    point_id = extract_number_from_text(title)
                    lat = float(row["Latitude"])
                    lon = float(row["Longitude"])
                    utm_x, utm_y, utm_zone, utm_band = latlon_to_utm(lat, lon, zone=target_zone)

                    date_created = row.get("Date Created") or ""
                    survey_date = ""
                    survey_time = ""
                    if date_created:
                        try:
                            dt = datetime.fromisoformat(date_created)
                            survey_date = dt.strftime("%Y-%m-%d")
                            survey_time = dt.strftime("%H:%M:%S")
                        except ValueError:
                            m = re.search(r"(\d{4}-\d{2}-\d{2})[T ](\d{2}:\d{2}:\d{2})", date_created)
                            if m:
                                survey_date, survey_time = m.group(1), m.group(2)

                    record = CoordinateRecord(
                        source_file=csv_file.name,
                        location=location,
                        day=day,
                        title=title,
                        point_id=point_id,
                        latitude=lat,
                        longitude=lon,
                        utm_x=utm_x,
                        utm_y=utm_y,
                        utm_zone=utm_zone,
                        utm_band=utm_band,
                        date_created=date_created,
                        survey_date=survey_date,
                        survey_time=survey_time,
                        elevation=row.get("Elevation", ""),
                    )

                    if point_id is not None:
                        records[(day, point_id)] = record
                except (KeyError, ValueError, TypeError):
                    continue

    return records


def format_exposure_time(exp_val: Any) -> Optional[str]:
    """Format EXIF ExposureTime (e.g. 0.00444 -> '1/225')."""
    if exp_val is None:
        return None
    try:
        val = float(exp_val)
        if val <= 0:
            return None
        if val >= 1.0:
            return f"{val:.1f}" if val % 1 else f"{int(val)}"
        reciprocal = round(1.0 / val)
        return f"1/{reciprocal}"
    except (ValueError, TypeError, ZeroDivisionError):
        return str(exp_val)


def parse_photo_filename(path: Path) -> Tuple[str, str, Optional[int], Optional[int], str]:
    """Parse photo name like 'D1E1 (1).JPG' or 'D1T1 (2).JPG'.

    Returns (day, transect, point_id, sequence, renamed_filename).
    E.g. 'D1E1 (1).JPG' -> day='D1', transect='T1', point_id=1, sequence=1, renamed='D1T1 (1).JPG'.
    """
    stem = path.stem
    ext = path.suffix.upper()

    day_match = re.search(r"D(\d+)", stem, re.IGNORECASE)
    day = f"D{day_match.group(1)}" if day_match else "D1"

    transect_match = re.search(r"[ET](\d+)", stem, re.IGNORECASE)
    transect = f"T{transect_match.group(1)}" if transect_match else "T1"

    seq_match = re.search(r"\((\d+)\)", stem)
    explicit_point = re.search(r"(?:P|TITIK|POINT)[ _-]?(\d+)", stem, re.IGNORECASE)

    if seq_match:
        point_id = int(seq_match.group(1))
        sequence = point_id
    elif explicit_point:
        point_id = int(explicit_point.group(1))
        sequence = point_id
    else:
        point_id = None
        sequence = None

    if point_id is not None:
        renamed_filename = f"{day}{transect} ({point_id}){ext}"
    else:
        renamed_filename = f"{day}{transect}_{stem}{ext}"

    return day, transect, point_id, sequence, renamed_filename


def extract_photo_metadata(path: Path) -> PhotoMetadata:
    """Extract EXIF and file metadata from a photo."""
    day, transect, point_id, sequence, renamed_filename = parse_photo_filename(path)
    file_size_bytes = os.path.getsize(path)
    file_size_str = f"{file_size_bytes / (1024 * 1024):.1f} MB"

    camera_model = None
    date_original = None
    time_original = None
    shutter_speed = None
    aperture = None
    iso = None
    white_balance = None

    try:
        with Image.open(path) as img:
            exif = img.getexif()
            if exif:
                camera_model = exif.get(ExifTags.Base.Model) or exif.get(0x0110)

            sub_exif = {}
            if hasattr(img, "_getexif") and img._getexif():
                sub_exif = img._getexif()

            dt_orig_raw = sub_exif.get(ExifTags.Base.DateTimeOriginal) or sub_exif.get(0x9003) or exif.get(ExifTags.Base.DateTime)
            if dt_orig_raw:
                if isinstance(dt_orig_raw, bytes):
                    dt_orig_raw = dt_orig_raw.decode("utf-8", errors="replace")
                dt_match = re.search(r"(\d{4})[:\-](\d{2})[:\-](\d{2})[ T](\d{2}:\d{2}:\d{2})", str(dt_orig_raw))
                if dt_match:
                    date_original = f"{dt_match.group(1)}-{dt_match.group(2)}-{dt_match.group(3)}"
                    time_original = dt_match.group(4)

            exp_time = sub_exif.get(ExifTags.Base.ExposureTime) or sub_exif.get(0x829A)
            if exp_time is not None:
                shutter_speed = format_exposure_time(exp_time)

            f_num = sub_exif.get(ExifTags.Base.FNumber) or sub_exif.get(0x829D)
            if f_num is not None:
                try:
                    aperture = f"{float(f_num):.1f}"
                except (ValueError, TypeError):
                    aperture = str(f_num)

            iso_val = sub_exif.get(ExifTags.Base.ISOSpeedRatings) or sub_exif.get(0x8827)
            if iso_val is not None:
                try:
                    iso = int(iso_val)
                except (ValueError, TypeError):
                    iso = None

            wb_val = sub_exif.get(ExifTags.Base.WhiteBalance) or sub_exif.get(0xA435)
            if wb_val is not None:
                white_balance = "Auto" if wb_val == 0 else "Manual"
    except Exception:
        pass

    return PhotoMetadata(
        source_path=path,
        original_filename=path.name,
        day=day,
        transect=transect,
        point_id=point_id,
        sequence=sequence,
        renamed_filename=renamed_filename,
        file_size_str=file_size_str,
        camera_model=camera_model or "V50 X Cube",
        date_original=date_original,
        time_original=time_original,
        shutter_speed=shutter_speed or "1/225",
        aperture=aperture or "2.8",
        iso=iso or 100,
        white_balance=white_balance or "Auto",
    )


def safe_copy_file(src: Path, dst: Path, copy_mode: str = "copy") -> None:
    """Copy or link a file, falling back when hardlinks are unsupported."""
    if copy_mode == "copy":
        try:
            shutil.copy2(src, dst)
        except OSError:
            shutil.copy(src, dst)
        return

    if copy_mode != "hardlink":
        raise ValueError("copy_mode must be 'copy' or 'hardlink'")

    try:
        dst.hardlink_to(src)
    except OSError:
        try:
            shutil.copy2(src, dst)
        except OSError:
            shutil.copy(src, dst)


def organize_photos(
    photos_dir: str | Path = "Photos",
    coords_dir: str | Path = "coords",
    output_dir: str | Path = "output",
    team: str = "A",
    copy_mode: str = "copy",
    url_prefix: str = "",
    utm_zone: Optional[int | str] = None,
    utm_format: str = "MGRS",
) -> Tuple[List[Dict[str, Any]], Path]:
    """Main function to organize photos, link coordinates, and export Excel manifest.

    Parameters:
    - photos_dir: Directory containing input photos (e.g. 'Photos')
    - coords_dir: Directory containing CSV coordinate files (e.g. 'coords')
    - output_dir: Directory where organized photos and Excel file will be placed
    - team: Team identifier for Column B (default: 'A')
    - copy_mode: 'copy' to duplicate, 'hardlink' to link without consuming space
    - url_prefix: Optional web URL prefix for Column R (URL).
    - utm_zone: Desired UTM Zone (e.g. 48 or 49 for Bangka Belitung, 51 for Sulawesi).
                If None or 'auto', automatically determined from coordinates.
    - utm_format: 'MGRS' (e.g. X_UTM48M, matches Example.xlsx) or 'HEMISPHERE' (e.g. X_UTM48S).

    Returns:
    - Tuple of (records list, Path to generated Excel file).
    """
    photos_path = Path(photos_dir)
    coords_path = Path(coords_dir)
    output_path = Path(output_dir)
    organized_photos_dir = output_path / "organized_photos"

    output_path.mkdir(parents=True, exist_ok=True)
    organized_photos_dir.mkdir(parents=True, exist_ok=True)

    target_zone_int = None
    if utm_zone and str(utm_zone).lower() != "auto":
        target_zone_int = int(re.search(r"\d+", str(utm_zone)).group())

    coord_lookup = load_coordinates(coords_path, target_zone=target_zone_int)

    # Determine representative UTM zone and band for the dataset
    if coord_lookup:
        zones = [pt.utm_zone for pt in coord_lookup.values()]
        bands = [pt.utm_band for pt in coord_lookup.values()]
        active_zone = target_zone_int or max(set(zones), key=zones.count)
        active_band = max(set(bands), key=bands.count)
    else:
        active_zone = target_zone_int or 51
        active_band = "M"

    col_x, col_y = get_utm_columns(active_zone, active_band, utm_format)
    columns = build_column_list(active_zone, active_band, utm_format)

    if not photos_path.exists():
        raise FileNotFoundError(
            f"Photos directory '{photos_path}' does not exist. "
            "Please check the path or ensure photos are uploaded/mounted."
        )

    photo_files = sorted([
        p for p in photos_path.iterdir()
        if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS
    ])

    records: List[Dict[str, Any]] = []

    for idx, photo_file in enumerate(photo_files, start=1):
        photo_meta = extract_photo_metadata(photo_file)

        coord = coord_lookup.get((photo_meta.day, photo_meta.point_id))

        point_num_str = str(photo_meta.point_id) if photo_meta.point_id is not None else str(idx)
        id_titik = f"{photo_meta.day}{photo_meta.transect}-{point_num_str}"

        transect_folder = organized_photos_dir / f"{photo_meta.day}{photo_meta.transect}"
        transect_folder.mkdir(parents=True, exist_ok=True)
        dest_photo_path = transect_folder / photo_meta.renamed_filename

        if not dest_photo_path.exists():
            safe_copy_file(photo_meta.source_path, dest_photo_path, copy_mode=copy_mode)

        rel_photo_path = os.path.relpath(dest_photo_path, output_path).replace("\\", "/")
        photo_url = f"{url_prefix.rstrip('/')}/{photo_meta.renamed_filename}" if url_prefix else rel_photo_path

        row_data = {
            # Columns A - P (Extracted from coordinate CSV)
            "ID": idx,
            "Team": team,
            "ID Transek": photo_meta.transect,
            "ID Titik": id_titik,
            "File name": photo_meta.renamed_filename,
            "Tanggal": coord.survey_date if coord else (photo_meta.date_original or ""),
            "Jam": coord.survey_time if coord else (photo_meta.time_original or ""),
            "Variasi Habitat": "?",
            "Kedalaman": coord.elevation if coord and coord.elevation else "?",
            "Range (m)": 0,
            "Azimuth": 90,
            "Latitude": coord.latitude if coord else None,
            "Longitude": coord.longitude if coord else None,
            "Lokasi": coord.location if coord else "Tandaigi",
            col_x: f"{coord.utm_x:.2f}" if coord else "",
            col_y: f"{coord.utm_y:.2f}" if coord else "",

            # Column Q & R (Photo image formula and link)
            "Photo": f'=IMAGE(R{idx + 1})',
            "URL": photo_url,

            # Column S (MUST remain empty for human annotator)
            "Annotation/Interpretation": "",

            # Columns T - AA (Extracted from photo image metadata)
            "FileSize": photo_meta.file_size_str,
            "Model": photo_meta.camera_model,
            "DateOriginal": photo_meta.date_original or "",
            "TimeOriginal": photo_meta.time_original or "",
            "ShutterSpeed": photo_meta.shutter_speed,
            "Aperture": photo_meta.aperture,
            "ISO": photo_meta.iso,
            "WhiteBalance": photo_meta.white_balance,
        }

        records.append(row_data)

    excel_file = output_path / "Survey_Organized.xlsx"
    export_excel_workbook(records, excel_file, columns=columns)

    return records, excel_file


def export_excel_workbook(
    records: List[Dict[str, Any]],
    output_file: Path,
    columns: Optional[List[str]] = None,
) -> None:
    """Generate the formatted Excel workbook matching Example.xlsx styling."""
    if columns is None:
        columns = build_column_list()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    header_font = Font(name="Aptos Narrow", size=11, bold=True, color="000000")
    data_font = Font(name="Aptos Narrow", size=11, bold=False, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

        if col_idx <= 16:  # Columns A - P: Yellow
            cell.fill = yellow_fill
        elif col_idx in (17, 18):  # Columns Q, R: Clean / white
            pass
        else:  # Columns S - AA: Accent fill
            cell.fill = accent_fill

    for row_idx, record in enumerate(records, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            val = record.get(col_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

            if col_idx <= 16:
                cell.fill = yellow_fill
                if col_name in ("ID", "Team", "ID Transek", "ID Titik", "Range (m)", "Azimuth"):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name in ("Latitude", "Longitude") or col_name.startswith("X_UTM") or col_name.startswith("Y_UTM"):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (17, 18):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.fill = accent_fill
                if col_name in ("Model", "Annotation/Interpretation"):
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    ws.freeze_panes = "A2"
    wb.save(output_file)


def main() -> None:
    parser = argparse.ArgumentParser(description="Organize survey photos and coordinates into Example.xlsx format.")
    parser.add_argument("--photos", default="Photos", help="Path to input photos folder (default: Photos)")
    parser.add_argument("--coords", default="coords", help="Path to coordinate CSV folder (default: coords)")
    parser.add_argument("--output", default="output", help="Output directory (default: output)")
    parser.add_argument("--team", default="A", help="Survey team code (default: A)")
    parser.add_argument("--copy-mode", choices=["copy", "hardlink"], default="copy", help="File copy method")
    parser.add_argument("--url-prefix", default="", help="Optional URL prefix for Google Drive or web hosting")
    parser.add_argument(
        "--utm-zone",
        default="auto",
        help="Target UTM Zone (e.g. 'auto', '48', '49', '51'). Useful for Bangka Belitung to unify projection.",
    )
    parser.add_argument(
        "--utm-format",
        choices=["MGRS", "HEMISPHERE"],
        default="MGRS",
        help="UTM Column suffix style: 'MGRS' (e.g. X_UTM48M) or 'HEMISPHERE' (e.g. X_UTM48S). Default: MGRS",
    )

    args = parser.parse_args()
    records, out_file = organize_photos(
        photos_dir=args.photos,
        coords_dir=args.coords,
        output_dir=args.output,
        team=args.team,
        copy_mode=args.copy_mode,
        url_prefix=args.url_prefix,
        utm_zone=args.utm_zone,
        utm_format=args.utm_format,
    )

    print(f"[SUCCESS] Processed and organized {len(records)} photos.")
    print(f"[SUCCESS] Excel workbook generated: {out_file.resolve()}")


if __name__ == "__main__":
    main()
